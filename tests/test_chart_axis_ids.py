import sys
import unittest
from pathlib import Path
from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
from chart_axis_ids import CHART_NS, axis_id_errors, normalize_chart_axis_ids


class ChartAxisIdsTest(unittest.TestCase):
    def chart(self, kind):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        if kind == XL_CHART_TYPE.XY_SCATTER:
            data = XyChartData()
            series = data.add_series('示例')
            series.add_data_point(1, -2)
            series.add_data_point(3, 4)
        else:
            data = CategoryChartData()
            data.categories = ['甲', '乙']
            data.add_series('示例', [-2, 4])
        return slide.shapes.add_chart(kind, Inches(1), Inches(1), Inches(6), Inches(4), data).chart._chartSpace

    def test_real_chart_templates_keep_series_and_cross_references(self):
        for kind in [XL_CHART_TYPE.BAR_CLUSTERED, XL_CHART_TYPE.COLUMN_CLUSTERED,
                     XL_CHART_TYPE.COLUMN_STACKED, XL_CHART_TYPE.LINE,
                     XL_CHART_TYPE.XY_SCATTER, XL_CHART_TYPE.DOUGHNUT]:
            with self.subTest(kind=kind):
                chart = self.chart(kind)
                ns = {'c': CHART_NS}
                before = [node.xml for node in chart.findall('.//c:ser', ns)]
                normalize_chart_axis_ids(chart)
                self.assertEqual(axis_id_errors(chart), [])
                self.assertEqual(before, [node.xml for node in chart.findall('.//c:ser', ns)])
                once = chart.xml
                normalize_chart_axis_ids(chart)
                self.assertEqual(once, chart.xml)

    def test_signed_reference_regression_and_valid_ids_unchanged(self):
        chart = self.chart(XL_CHART_TYPE.COLUMN_CLUSTERED)
        self.assertTrue(axis_id_errors(chart))  # Existing public template fails strict readers.
        normalize_chart_axis_ids(chart)
        before = chart.xml
        normalize_chart_axis_ids(chart)
        self.assertEqual(before, chart.xml)

    def test_dangling_axis_reference_fails(self):
        chart = self.chart(XL_CHART_TYPE.COLUMN_CLUSTERED)
        normalize_chart_axis_ids(chart)
        chart.find('.//c:crossAx', {'c': CHART_NS}).set('val', '12345')
        self.assertTrue(axis_id_errors(chart))
        with self.assertRaises(ValueError):
            normalize_chart_axis_ids(chart)

    def test_out_of_range_and_collision_fail(self):
        base = self.chart(XL_CHART_TYPE.COLUMN_CLUSTERED)
        node = base.find('.//c:axId', {'c': CHART_NS})
        node.set('val', str(2**32))
        with self.assertRaises(ValueError):
            normalize_chart_axis_ids(base)
        base = self.chart(XL_CHART_TYPE.COLUMN_CLUSTERED)
        nodes = base.findall('.//c:axId', {'c': CHART_NS})
        nodes[0].set('val', '-1')
        nodes[1].set('val', str(2**32-1))
        with self.assertRaises(ValueError):
            normalize_chart_axis_ids(base)


if __name__ == '__main__':
    unittest.main()
